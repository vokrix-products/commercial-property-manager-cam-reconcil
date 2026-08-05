import os
import json
import io
import pdfplumber
import openpyxl
from openai import OpenAI
from datetime import datetime

CHARGE_THRESHOLD = 5000.0

def process_file(file_bytes: bytes) -> list[dict]:
    text = ""
    # Try PDF
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        if text.strip():
            return _extract_records(text)
    except Exception:
        pass
    # Try Excel
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                text += row_text + "\n"
        if text.strip():
            return _extract_records(text)
    except Exception:
        pass
    # Fallback: plain text/CSV
    try:
        text = file_bytes.decode("utf-8", errors="ignore")
        if text.strip():
            return _extract_records(text)
        else:
            return []
    except Exception:
        return []

def _extract_records(text: str) -> list[dict]:
    client = OpenAI(
        api_key=os.environ["DEEPSEEK_API_KEY"],
        base_url="https://api.deepseek.com"
    )
    prompt = f"""
You are a data extraction assistant specialized in commercial property management CAM reconciliations.
Extract tenant charge records from the following document text.
For each tenant, return a JSON object with the following fields:
- "tenant_name": the name of the tenant (primary entity being charged)
- "charge_amount": the total charge amount as a float number (in dollars)
- "property": the property or building name if mentioned, otherwise null
- "due_date": the due date of the charge in ISO-8601 format (YYYY-MM-DD) if available, otherwise null

Return the result as a JSON array of objects, even if only one record.
Do not include any markdown or additional text. Just the raw JSON array.
Document text:
{text}
"""
    response = client.chat.completions.create(
        model="deepseek-chat",
        messages=[{"role": "user", "content": prompt}],
        temperature=0,
        max_tokens=2000
    )
    content = response.choices[0].message.content.strip()
    try:
        data = json.loads(content)
    except json.JSONDecodeError:
        import re
        json_match = re.search(r'\[.*\]', content, re.DOTALL)
        if json_match:
            data = json.loads(json_match.group())
        else:
            return []
    records = []
    for item in data:
        tenant_name = item.get("tenant_name", "Unknown")
        charge_amount = item.get("charge_amount", 0.0)
        status = "above_threshold:critical" if charge_amount > CHARGE_THRESHOLD else "within_threshold:good"
        due_date = item.get("due_date")
        if due_date and isinstance(due_date, str):
            try:
                datetime.fromisoformat(due_date)
            except ValueError:
                due_date = None
        elif not due_date:
            due_date = None
        records.append({
            "title": tenant_name,
            "status": status,
            "details": {
                "charge_amount": charge_amount,
                "property": item.get("property", ""),
            },
            "due_date": due_date
        })
    return records
